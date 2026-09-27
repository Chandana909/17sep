"""Source readers: turn an extract file into raw rows (column name -> value).

Readers only read. Mapping, typing and validation happen in `ingest.py`, so adding a source
format is one function registered in `READERS` (for example a read-only warehouse query), and
nothing downstream changes. CSV and JSON Lines need no extra packages. Parquet needs
`pyarrow` and Excel needs `openpyxl`; both are imported only when used.
"""

from __future__ import annotations

import csv
import json
from collections.abc import Callable, Iterator, Mapping
from datetime import date, datetime
from pathlib import Path
from typing import Any

from asas.core.errors import DataContractError

Row = dict[str, object]
Reader = Callable[[Path, Mapping[str, str]], Iterator[Row]]


def read_csv(path: Path, options: Mapping[str, str]) -> Iterator[Row]:
    encoding = options.get("encoding", "utf-8-sig")
    delimiter = options.get("delimiter", ",")
    with open(path, encoding=encoding, newline="") as fh:
        for row in csv.DictReader(fh, delimiter=delimiter):
            if None in row:
                raise DataContractError(f"{path.name}: a row has more cells than the header")
            if any(v not in ("", None) for v in row.values()):
                yield dict(row)


def read_jsonl(path: Path, options: Mapping[str, str]) -> Iterator[Row]:
    encoding = options.get("encoding", "utf-8")
    with open(path, encoding=encoding) as fh:
        for number, line in enumerate(fh, start=1):
            if not line.strip():
                continue
            try:
                value = json.loads(line)
            except json.JSONDecodeError as exc:
                raise DataContractError(f"{path.name}:{number}: not valid JSON") from exc
            if not isinstance(value, dict):
                raise DataContractError(f"{path.name}:{number}: each line must be a JSON object")
            yield value


def read_parquet(path: Path, options: Mapping[str, str]) -> Iterator[Row]:
    try:
        import pyarrow.parquet as pq
    except ImportError as exc:  # pragma: no cover - depends on the environment
        raise DataContractError("reading parquet needs the 'pyarrow' package") from exc
    table: Any = pq.read_table(path)
    yield from table.to_pylist()


def _cell(value: object) -> object:
    if isinstance(value, datetime):
        return value
    if isinstance(value, date):
        return datetime(value.year, value.month, value.day)
    return value


def read_xlsx(path: Path, options: Mapping[str, str]) -> Iterator[Row]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover - depends on the environment
        raise DataContractError("reading xlsx needs the 'openpyxl' package") from exc
    book: Any = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = book[options["sheet"]] if "sheet" in options else book.active
        rows = sheet.iter_rows(values_only=True)
        header = [str(h).strip() if h is not None else "" for h in next(rows, ())]
        for values in rows:
            if all(v is None for v in values):
                continue
            yield {h: _cell(v) for h, v in zip(header, values, strict=False) if h}
    finally:
        book.close()


READERS: dict[str, Reader] = {
    "csv": read_csv,
    "jsonl": read_jsonl,
    "parquet": read_parquet,
    "xlsx": read_xlsx,
}
EXTENSIONS = {
    ".csv": "csv",
    ".txt": "csv",
    ".tsv": "csv",
    ".jsonl": "jsonl",
    ".ndjson": "jsonl",
    ".parquet": "parquet",
    ".xlsx": "xlsx",
}


def reader_name(path: Path, name: str | None = None) -> str:
    chosen = name or EXTENSIONS.get(path.suffix.lower())
    if chosen is None or chosen not in READERS:
        known = ", ".join(sorted(READERS))
        raise DataContractError(
            f"{path.name}: no reader for this file; set `reader` in the mapping to one of {known}"
        )
    return chosen


def read_rows(
    path: Path, name: str | None = None, options: Mapping[str, str] | None = None
) -> list[Row]:
    if not path.is_file():
        raise DataContractError(f"source file not found: {path}")
    opts = dict(options or {})
    chosen = reader_name(path, name)
    if chosen == "csv" and path.suffix.lower() == ".tsv":
        opts.setdefault("delimiter", "\t")
    return list(READERS[chosen](path, opts))
